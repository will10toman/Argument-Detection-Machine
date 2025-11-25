from fastapi import UploadFile
import PyPDF2
from docx import Document
import aiofiles
import openpyxl
import os
import uuid

async def extract_text_from_file(file: UploadFile):
    temp_dir = "storage/temp"
    os.makedirs(temp_dir, exist_ok=True)

    uid = str(uuid.uuid4())
    ext = file.filename.split(".")[-1]
    path = f"{temp_dir}/{uid}.{ext}"

    async with aiofiles.open(path, "wb") as f:
        await f.write(await file.read())

    try:
        if ext == "pdf":
            text = extract_from_pdf(path)
        elif ext == "docx":
            text = extract_from_docx(path)
        elif ext == "txt":
            async with aiofiles.open(path, "r", encoding="utf-8") as f:
                text = await f.read()
        elif ext == "xlsx":
            text = extract_from_xlsx(path)
        else:
            raise ValueError("Unsupported file type")

        return text
    finally:
        os.remove(path)

def extract_from_pdf(path):
    with open(path, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        return "\n".join([p.extract_text() for p in reader.pages])

def extract_from_docx(path):
    doc = Document(path)
    return "\n".join([p.text for p in doc.paragraphs])

def extract_from_xlsx(path):
    wb = openpyxl.load_workbook(path)
    text = ""
    for sheet in wb:
        for row in sheet.iter_rows():
            text += " ".join([str(cell.value) for cell in row if cell.value]) + "\n"
    return text
